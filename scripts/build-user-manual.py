from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "documentacion"
OUTPUT = OUTPUT_DIR / "Manual_de_Usuario_Abicorp_ERP_v0.1.docx"
LOGO = ROOT / "public" / "assets" / "abicorp-logo.png"

NAVY = "064B79"
TEAL = "009EA5"
FOREST = "123D2C"
LIME = "C8F36B"
INK = "17231E"
MUTED = "68736D"
LINE = "DCE5E0"
PALE = "EEF6F4"
PALE_BLUE = "EAF2F8"
PALE_GOLD = "FFF7E8"
RED = "A33F3B"
WHITE = "FFFFFF"
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb(value):
    return RGBColor.from_string(value)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA):
    if sum(widths) != USABLE_DXA:
        raise ValueError(f"Las columnas deben sumar {USABLE_DXA} DXA: {widths}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Calibri", size=None, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.widow_control = True

    for style_name, size, color, before, after in (
        ("Heading 1", 16, NAVY, 18, 10),
        ("Heading 2", 13, TEAL, 14, 7),
        ("Heading 3", 12, NAVY, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr, fld_sep, cached, fld_char_2])
    set_font(run, size=9, color=MUTED)


def paragraph_top_border(paragraph, color=LINE, size=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), str(size))
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), color)
    p_bdr.append(top)


def configure_page(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    first_header = section.first_page_header
    first_header.paragraphs[0].text = ""

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    logo_shape = run.add_picture(str(LOGO), width=Inches(0.24))
    logo_shape._inline.docPr.set("descr", "Logotipo de Abicorp ERP")
    logo_shape._inline.docPr.set("title", "Abicorp ERP")
    text = p.add_run("   ABICORP ERP  |  MANUAL DE USUARIO")
    set_font(text, size=8.5, color=MUTED, bold=True)
    paragraph_top_border(p, color=LINE, size=4)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph_top_border(fp, color=LINE, size=4)
    r = fp.add_run("M4rk3z Solutions  •  Versión 0.1  •  Página ")
    set_font(r, size=8.5, color=MUTED)
    add_page_field(fp)

    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    ffp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = ffp.add_run("Todos los derechos reservados M4rk3z Solutions")
    set_font(r, size=8.5, color=MUTED)


def make_numbering(doc, kind="decimal"):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_number(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def add_bullets(doc, items):
    num_id = make_numbering(doc, "bullet")
    for item in items:
        p = doc.add_paragraph()
        apply_number(p, num_id)
        if isinstance(item, tuple):
            label, detail = item
            r = p.add_run(label + ": ")
            set_font(r, bold=True)
            r = p.add_run(detail)
            set_font(r)
        else:
            r = p.add_run(item)
            set_font(r)


def add_steps(doc, items):
    num_id = make_numbering(doc, "decimal")
    for label, detail in items:
        p = doc.add_paragraph()
        apply_number(p, num_id)
        r = p.add_run(label + ". ")
        set_font(r, bold=True, color=NAVY)
        r = p.add_run(detail)
        set_font(r)


def add_note(doc, title, text, kind="info"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.05)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_GOLD if kind == "warning" else PALE_BLUE)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "D39B3A" if kind == "warning" else TEAL)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    r = p.add_run(title.upper() + "\n")
    set_font(r, size=9, color="8A5A00" if kind == "warning" else TEAL, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, FOREST)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_font(r, size=9.5, color=WHITE, bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            if row_idx % 2:
                set_cell_shading(cells[idx], "F7F9F8")
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_font(r, size=9.5, color=INK)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_section_title(doc, title, intro=None, new_page=True):
    if new_page:
        doc.add_page_break()
    doc.add_heading(title, level=1)
    if intro:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        r = p.add_run(intro)
        set_font(r, size=11.5, color=MUTED, italic=True)


def add_module(doc, module):
    add_section_title(doc, module["title"], module["purpose"])
    doc.add_heading("Antes de comenzar", level=2)
    add_note(doc, "Requisitos", module["before"])
    doc.add_heading("Procedimiento recomendado", level=2)
    add_steps(doc, module["steps"])
    doc.add_heading("Controles y resultados", level=2)
    add_bullets(doc, module["controls"])
    doc.add_heading("Conexiones con otros módulos", level=2)
    p = doc.add_paragraph(module["connections"])
    p.paragraph_format.space_after = Pt(6)


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    doc.core_properties.title = "Manual de Usuario - Abicorp ERP"
    doc.core_properties.subject = "Guía operativa del ERP modular Abicorp"
    doc.core_properties.author = "M4rk3z Solutions"
    doc.core_properties.keywords = "Abicorp, ERP, manual de usuario, operación"
    doc.core_properties.comments = "Documento de operación para la versión 0.1"

    # Portada editorial
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(62)
    run = p.add_run()
    logo_shape = run.add_picture(str(LOGO), width=Inches(1.35))
    logo_shape._inline.docPr.set("descr", "Logotipo de Abicorp ERP")
    logo_shape._inline.docPr.set("title", "Abicorp ERP")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(22)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("MANUAL DE USUARIO")
    set_font(r, size=11, color=TEAL, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Abicorp ERP")
    set_font(r, size=31, color=NAVY, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("Tu operación completa, en un solo lugar")
    set_font(r, size=15, color=TEAL, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Inches(0.65)
    p.paragraph_format.right_indent = Inches(0.65)
    r = p.add_run(
        "Guía práctica para operar el núcleo, los datos maestros, almacén, compras, "
        "ventas, producción, calidad, mantenimiento, logística, finanzas y aprobaciones."
    )
    set_font(r, size=11.5, color=MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    r = p.add_run("VERSIÓN 0.1  |  JULIO DE 2026")
    set_font(r, size=10, color=FOREST, bold=True)

    # Índice
    add_section_title(doc, "Contenido", "Ruta de consulta rápida del manual.", new_page=True)
    contents = [
        "1. Acerca del sistema",
        "2. Puesta en marcha y acceso",
        "3. Navegación de la interfaz",
        "4. Roles, niveles y permisos",
        "5. Dashboards y atajos",
        "6. Núcleo del sistema",
        "7. Datos maestros",
        "8. Almacén e inventario",
        "9. Compras",
        "10. Ventas",
        "11. Producción",
        "12. Calidad",
        "13. Mantenimiento",
        "14. Logística y embarques",
        "15. Finanzas",
        "16. Tareas y aprobaciones",
        "17. Flujo integral entre módulos",
        "18. Buenas prácticas operativas",
        "19. Solución de problemas",
        "20. Respaldo y administración local",
        "21. Referencia rápida",
    ]
    add_bullets(doc, contents)
    add_note(
        doc,
        "Alcance",
        "Las opciones visibles dependen del rol del usuario. Si un botón no aparece, "
        "confirme primero que la cuenta tenga el permiso necesario.",
    )

    # 1
    add_section_title(
        doc,
        "1. Acerca del sistema",
        "Abicorp ERP es una aplicación web modular con servidor y base SQLite locales.",
    )
    doc.add_heading("Qué resuelve", level=2)
    add_bullets(
        doc,
        [
            ("Una sola operación", "los documentos se conectan desde la solicitud hasta el cierre."),
            ("Trazabilidad", "folios, estados, responsables y movimientos quedan relacionados."),
            ("Datos locales", "la información reside en la base SQLite del equipo o servidor local."),
            ("Control por rol", "cada usuario consulta, opera, valida o administra según su nivel."),
            ("Folios automáticos", "el sistema asigna el consecutivo al guardar cada registro."),
        ],
    )
    doc.add_heading("Mapa de módulos", level=2)
    add_table(
        doc,
        ["Área", "Uso principal", "Resultado"],
        [
            ("Núcleo", "Usuarios, permisos, catálogos y configuración", "Gobierno del sistema"),
            ("Datos maestros", "Artículos, clientes, proveedores, empleados y recursos", "Información base confiable"),
            ("Almacén", "Existencias, cargas, retiros, movimientos y conteos", "Inventario controlado"),
            ("Compras", "Solicitud, proveedor, orden, recepción y factura", "Adquisiciones trazables"),
            ("Ventas", "Prospecto, cotización, pedido, entrega y factura", "Seguimiento comercial"),
            ("Producción", "Materiales, rutas, planeación, órdenes y ejecución", "Avance de fabricación"),
            ("Calidad", "Planes, inspecciones, no conformidades y liberación", "Producto verificado"),
            ("Mantenimiento", "Equipos, prevención, solicitudes, órdenes y paros", "Disponibilidad de activos"),
            ("Logística", "Surtido, empaque, ruta, evidencia y entrega", "Distribución controlada"),
            ("Finanzas", "Cuentas, pagos, cobros, presupuestos y conciliación", "Control financiero"),
            ("Tareas", "Asignaciones, aprobaciones, comentarios y decisiones", "Seguimiento de pendientes"),
        ],
        [1800, 4200, 3360],
    )

    # 2
    add_section_title(
        doc,
        "2. Puesta en marcha y acceso",
        "El navegador muestra la interfaz; el servidor local atiende la autenticación y la base de datos.",
    )
    doc.add_heading("Arrancar el sistema", level=2)
    add_steps(
        doc,
        [
            ("Abra PowerShell", "sitúese en la carpeta del proyecto Abicorp."),
            ("Inicie el servidor", "ejecute npm start y mantenga esa ventana abierta."),
            ("Abra la aplicación", "entre a http://127.0.0.1:5050 desde el navegador."),
            ("Inicie sesión", "capture el usuario y la contraseña proporcionados por el administrador."),
            ("Cambie la contraseña", "si el sistema lo solicita en el primer acceso, establezca una clave nueva."),
        ],
    )
    add_note(
        doc,
        "Live Server",
        "Puede mostrar public/index.html con Live Server, pero npm start debe seguir ejecutándose. "
        "De lo contrario aparecerá “Failed to fetch” porque no existe conexión con el servidor y SQLite.",
        "warning",
    )
    doc.add_heading("Inicio y cierre de sesión", level=2)
    add_bullets(
        doc,
        [
            "El usuario activo aparece en la parte superior derecha.",
            "La fecha y la hora se actualizan con la zona horaria configurada.",
            "Use el icono de salida junto al nombre para cerrar la sesión de forma segura.",
            "No comparta contraseñas ni deje una sesión abierta en un equipo compartido.",
        ],
    )

    # 3
    add_section_title(
        doc,
        "3. Navegación de la interfaz",
        "La interfaz está diseñada para concentrar cada área en una pantalla operativa.",
    )
    add_table(
        doc,
        ["Elemento", "Ubicación", "Uso"],
        [
            ("Menú lateral", "Lado izquierdo", "Abre Dashboards y los módulos operativos."),
            ("Grupos desplegables", "Núcleo y Tareas", "Muestran varias pantallas relacionadas."),
            ("Encabezado", "Parte superior", "Título, ruta, fecha, hora, usuario, salida y notificaciones."),
            ("Área de trabajo", "Centro", "Indicadores, registros, acciones y formularios."),
            ("Ventanas modales", "Sobre la pantalla", "Creación, edición, detalle o ejecución de una acción."),
            ("Mensajes emergentes", "Esquina de la pantalla", "Confirman el resultado o muestran un error."),
        ],
        [1800, 2200, 5360],
    )
    doc.add_heading("Reglas generales de captura", level=2)
    add_bullets(
        doc,
        [
            ("Folio", "no se captura; se asigna al guardar según el tipo y el consecutivo."),
            ("Campos obligatorios", "están marcados por el formulario y deben completarse para continuar."),
            ("Listas", "solo muestran registros activos y permitidos para la operación."),
            ("Estados", "determinan qué acciones están disponibles en cada documento."),
            ("Cancelar", "cierra el formulario sin aplicar cambios."),
        ],
    )

    # 4
    add_section_title(
        doc,
        "4. Roles, niveles y permisos",
        "Los niveles definen el alcance máximo; los permisos habilitan acciones específicas.",
    )
    add_table(
        doc,
        ["Nivel", "Nombre", "Alcance habitual"],
        [
            ("1", "Consultar", "Ver información, indicadores y detalles."),
            ("2", "Crear y editar", "Registrar y modificar operaciones permitidas."),
            ("3", "Validar y aprobar", "Aprobar, rechazar, liberar o confirmar documentos."),
            ("4", "Administrar", "Configurar, asignar permisos y controlar el sistema."),
        ],
        [1200, 2600, 5560],
    )
    doc.add_heading("Asignar un rol", level=2)
    add_steps(
        doc,
        [
            ("Abra Núcleo", "seleccione Usuarios para crear o localizar la cuenta."),
            ("Revise el área", "asigne la unidad organizacional correspondiente."),
            ("Seleccione el rol", "elija el rol que refleje las responsabilidades reales."),
            ("Verifique permisos", "en Roles y permisos confirme las capacidades habilitadas."),
            ("Pruebe el acceso", "valide con la cuenta que solo aparezcan las opciones necesarias."),
        ],
    )
    add_note(
        doc,
        "Principio de mínimo privilegio",
        "Asigne únicamente los permisos necesarios para la función. Las aprobaciones y cierres deben "
        "reservarse para responsables autorizados.",
        "warning",
    )

    # 5
    add_section_title(
        doc,
        "5. Dashboards y atajos",
        "Dashboards es el espacio personal de acceso rápido de cada usuario.",
    )
    add_steps(
        doc,
        [
            ("Abra Dashboards", "seleccione el primer acceso del menú lateral."),
            ("Agregue un atajo", "use “Agregar atajo” junto al contador de Mis atajos."),
            ("Seleccione módulos", "marque uno o varios accesos disponibles para su rol."),
            ("Abra un módulo", "seleccione la tarjeta para ir directamente al área."),
            ("Quite un atajo", "use la × de la tarjeta; el módulo no se elimina del sistema."),
        ],
    )
    add_note(
        doc,
        "Preferencia local",
        "Los atajos se guardan para el usuario en el navegador actual. Limpiar los datos del navegador "
        "puede restablecer esa selección.",
    )

    # 6
    add_section_title(
        doc,
        "6. Núcleo del sistema",
        "El núcleo administra usuarios, reglas, catálogos, documentos y supervisión.",
    )
    doc.add_heading("Funciones principales", level=2)
    add_table(
        doc,
        ["Opción", "Qué administra", "Acción frecuente"],
        [
            ("Usuarios", "Cuentas, rol, área, estado y bloqueo", "Crear, editar o activar cuentas"),
            ("Roles y permisos", "Nivel y permisos por módulo", "Definir quién consulta, opera o aprueba"),
            ("Catálogos generales", "Empresas, sucursales, áreas, almacenes, unidades, monedas y estados", "Mantener listas base"),
            ("Folios", "Prefijos, consecutivos y reinicios", "Consultar o ajustar series autorizadas"),
            ("Archivos", "Documentos locales de hasta 8 MB", "Subir, consultar y descargar"),
            ("Notificaciones", "Avisos individuales o generales", "Enviar y marcar como leído"),
            ("Bitácora", "Eventos y movimientos del sistema", "Investigar quién hizo qué y cuándo"),
            ("Configuración", "Zona horaria y parámetros generales", "Guardar ajustes y revisar supervisión"),
        ],
        [1900, 4000, 3460],
    )
    doc.add_heading("Supervisión del sistema", level=2)
    add_bullets(
        doc,
        [
            "Abra Configuración y cambie a la pestaña Supervisión del sistema.",
            "Revise usuarios activos, áreas, sesiones y eventos recientes.",
            "Use la bitácora cuando necesite el detalle completo de una operación.",
            "Confirme que el indicador SQLite local conectado permanezca activo.",
        ],
    )

    modules = [
        {
            "title": "7. Datos maestros",
            "purpose": "Centraliza la información reutilizada por todos los documentos del ERP.",
            "before": "Registre primero empresas, sucursales, áreas, unidades de medida y monedas en el núcleo.",
            "steps": [
                ("Abra Datos maestros", "elija la categoría interna que desea consultar."),
                ("Seleccione Nuevo", "abra el formulario del registro correspondiente."),
                ("Capture los datos", "complete nombre, clasificación, moneda, unidad y condiciones aplicables."),
                ("Guarde", "el sistema asignará automáticamente el código o folio."),
                ("Verifique", "use búsqueda y filtros para confirmar que el registro quedó activo."),
            ],
            "controls": [
                ("Artículos", "productos terminados, materias primas, consumibles, herramientas, servicios y activos."),
                ("Clientes", "datos fiscales, contacto, moneda, crédito y límite."),
                ("Proveedores", "condiciones comerciales y tiempo de entrega."),
                ("Empleados", "número, área, puesto y estado."),
                ("Recursos", "máquinas, centros de trabajo, herramientas y personal productivo."),
                ("Precios y costos", "listas vigentes y partidas por artículo."),
            ],
            "connections": "Los artículos alimentan compras, ventas, inventario, producción y mantenimiento. "
            "Clientes y proveedores habilitan los documentos comerciales; empleados y recursos se usan en ejecución y asignaciones.",
        },
        {
            "title": "8. Almacén e inventario",
            "purpose": "Muestra qué existe en cada almacén y permite cargar, retirar o mover artículos desde una sola vista.",
            "before": "Debe existir al menos un almacén activo, un artículo inventariable y su unidad de medida.",
            "steps": [
                ("Seleccione un almacén", "abra su tarjeta para consultar existencias y ubicaciones."),
                ("Busque el artículo", "filtre por código o nombre antes de operar."),
                ("Elija la acción", "use Cargar, Retirar o Mover según el movimiento físico."),
                ("Capture cantidades", "indique almacén, ubicación, lote o serie cuando corresponda."),
                ("Confirme", "revise el folio automático y la existencia actualizada."),
            ],
            "controls": [
                ("Disponible", "existencia física menos la cantidad reservada."),
                ("Reservas", "separan cantidad sin retirarla; pueden liberarse o consumirse."),
                ("Lotes y series", "mantienen trazabilidad de partidas y unidades."),
                ("Ubicaciones", "organizan zona, pasillo, rack, nivel y contenedor."),
                ("Conteos", "comparan el saldo del sistema contra el conteo observado."),
                ("Saldo negativo", "el sistema bloquea salidas superiores a la cantidad disponible."),
            ],
            "connections": "Las recepciones de Compras generan entradas; Ventas y Logística generan salidas; "
            "Producción consume materiales y registra producto terminado; Mantenimiento consume refacciones.",
        },
        {
            "title": "9. Compras",
            "purpose": "Controla la adquisición desde la necesidad interna hasta la recepción y la cuenta por pagar.",
            "before": "Registre artículos comprables, proveedores, almacenes, moneda y, si aplica, centros de costo.",
            "steps": [
                ("Cree la solicitud", "capture necesidad, fecha requerida, prioridad, centro de costo y partidas."),
                ("Compare proveedores", "registre ofertas y seleccione la alternativa autorizada."),
                ("Genere la orden", "tome la solicitud, proveedor y precios; después apruebe la compra."),
                ("Reciba el material", "capture cantidades recibidas, almacén, ubicación y referencia del proveedor."),
                ("Registre la factura", "vincule la orden y confirme subtotal, impuestos, total y vencimiento."),
                ("Gestione excepciones", "use Devolución cuando material recibido deba regresar al proveedor."),
            ],
            "controls": [
                "La recepción actualiza inventario únicamente para artículos con control de existencia.",
                "Las recepciones parciales conservan la cantidad pendiente de cada partida.",
                "La devolución descuenta del almacén las cantidades regresadas.",
                "La factura del proveedor crea una cuenta por pagar enlazada en Finanzas.",
                "Las aprobaciones dependen del permiso y del estado del documento.",
            ],
            "connections": "Inventario puede originar la necesidad por material faltante. Compras abastece almacén y "
            "Finanzas recibe automáticamente el compromiso de pago.",
        },
        {
            "title": "10. Ventas",
            "purpose": "Concentra prospectos, cotizaciones, pedidos, entregas, devoluciones y facturación.",
            "before": "Registre clientes, artículos vendibles, listas de precio, moneda e inventario disponible.",
            "steps": [
                ("Registre el prospecto", "capture empresa, contacto, origen, etapa y valor estimado."),
                ("Prepare la cotización", "seleccione cliente y agregue partidas con precio, descuento e impuesto."),
                ("Acepte la propuesta", "cambie el estado cuando el cliente confirme."),
                ("Cree el pedido", "use Crear pedido desde la cotización; los datos y partidas se precargan."),
                ("Confirme el pedido", "revise fechas, crédito y entrega esperada antes de continuar."),
                ("Complete el ciclo", "registre entrega, devolución o cancelación y genere la factura cuando corresponda."),
            ],
            "controls": [
                "Cada documento conserva el vínculo con su documento de origen.",
                "Los artículos del selector deben estar activos y habilitados para venta.",
                "El pedido generado desde una cotización no requiere recapturar cliente ni partidas.",
                "Entregas y devoluciones actualizan el inventario.",
                "La factura comercial genera seguimiento de cobro en Finanzas.",
            ],
            "connections": "Un pedido confirmado puede iniciar Producción o pasar a Logística si ya existe producto disponible. "
            "Facturación alimenta cuentas por cobrar.",
        },
        {
            "title": "11. Producción",
            "purpose": "Controla ingeniería, planeación, órdenes, materiales, operaciones y avance de fabricación.",
            "before": "Registre productos, materias primas, recursos, centros de trabajo, almacenes y rutas.",
            "steps": [
                ("Defina materiales", "cree la lista de materiales y su versión para el producto."),
                ("Defina el proceso", "registre ruta, operaciones, secuencia, recursos y tiempos estándar."),
                ("Planee la demanda", "capture el requerimiento y revise material y capacidad."),
                ("Cree la orden", "seleccione producto, cantidad, fechas y documentos relacionados."),
                ("Ejecute", "inicie, pause, reanude o termine operaciones; registre cantidad, rechazo y tiempos."),
                ("Cierre", "confirme consumos, entrada de producto terminado, subproductos y desperdicios."),
            ],
            "controls": [
                "La orden conserva materiales y operaciones requeridas.",
                "La revisión de inventario identifica disponibilidad o faltantes.",
                "Los faltantes pueden derivarse a Compras.",
                "Los consumos reducen materia prima y la entrada aumenta producto terminado.",
                "El producto final puede quedar pendiente de liberación por Calidad.",
            ],
            "connections": "Recibe demanda de Ventas, consulta y mueve Inventario, solicita faltantes a Compras, usa recursos "
            "de Datos maestros y envía el resultado a Calidad.",
        },
        {
            "title": "12. Calidad",
            "purpose": "Administra planes de inspección, resultados, no conformidades, liberaciones y correcciones.",
            "before": "Registre artículos, planes de inspección y criterios aplicables al origen de la revisión.",
            "steps": [
                ("Cree el plan", "defina el tipo de inspección, características, método y criterio."),
                ("Abra una inspección", "seleccione recepción, proceso, producto final o registro aplicable."),
                ("Capture resultados", "registre mediciones, observaciones y evidencia requerida."),
                ("Evalúe", "marque el resultado como aprobado o rechazado conforme al plan."),
                ("Trate el rechazo", "genere producto no conforme y una acción correctiva si procede."),
                ("Libere", "autorice el producto aprobado para continuar al almacén o embarque."),
            ],
            "controls": [
                "Recepción: valida material comprado antes de uso.",
                "Proceso: verifica el producto durante fabricación.",
                "Final: confirma cumplimiento antes de liberar.",
                "No conformidad: bloquea o dirige a reproceso el producto rechazado.",
                "Acción correctiva: documenta causa, responsable y seguimiento.",
            ],
            "connections": "Calidad recibe materiales de Compras y producto de Producción. Un resultado aprobado libera el flujo "
            "hacia Inventario y Logística; un rechazo activa bloqueo, devolución o reproceso.",
        },
        {
            "title": "13. Mantenimiento",
            "purpose": "Reúne equipos, prevención, solicitudes, órdenes, refacciones, paros e historial en una sola pantalla.",
            "before": "Registre el equipo, su ubicación, criticidad, medidor y, si aplica, recursos y refacciones.",
            "steps": [
                ("Registre el equipo", "capture ficha técnica, ubicación, criticidad y lectura inicial."),
                ("Programe prevención", "cree un plan por calendario o por lectura del medidor."),
                ("Reporte la necesidad", "abra una solicitud con falla, prioridad y condición de paro."),
                ("Cree la orden", "vincule solicitud o plan, asigne técnico, fechas y refacciones."),
                ("Ejecute el trabajo", "apruebe, inicie, pause, reanude, consuma refacciones y registre paro."),
                ("Termine y cierre", "documente resultado, actualice el equipo y consulte el historial."),
            ],
            "controls": [
                "Los equipos pueden estar operativos, en mantenimiento, detenidos o retirados.",
                "Los planes vencidos se identifican para generar una orden preventiva.",
                "Las refacciones utilizadas generan una salida de inventario.",
                "Los paros calculan el tiempo desde el inicio hasta la reanudación.",
                "El historial conserva solicitudes, órdenes, consumos, paros y cierres.",
            ],
            "connections": "Usa artículos de Datos maestros e Inventario para refacciones, empleados y recursos para asignación, "
            "y registra eventos auditables en el núcleo.",
        },
        {
            "title": "14. Logística y embarques",
            "purpose": "Controla preparación, surtido, empaque, ruta, evidencia y entrega al cliente.",
            "before": "Debe existir un pedido confirmado, almacén, existencia disponible y, para el despacho, una ruta o transportista.",
            "steps": [
                ("Prepare el pedido", "convierta el pedido confirmado en un embarque; las partidas se copian automáticamente."),
                ("Inicie el surtido", "registre ubicación, lote y cantidades tomadas."),
                ("Termine picking", "confirme que todas las partidas estén completas."),
                ("Empaque", "capture bultos, peso y referencia de rastreo."),
                ("Asigne ruta", "seleccione trayecto, conductor, vehículo y transportista."),
                ("Despache y entregue", "registre evidencia y confirme la recepción del cliente."),
            ],
            "controls": [
                "Estados habituales: preparación, picking, surtido, empacado, en tránsito y entregado.",
                "La evidencia puede contener fotografía, firma, archivo o nota.",
                "La confirmación final genera la entrega relacionada en Ventas.",
                "La salida de inventario se aplica conforme avanza la entrega.",
                "El pedido, almacén, ruta, evidencia y cliente permanecen vinculados.",
            ],
            "connections": "Recibe pedidos de Ventas y existencias de Almacén. Al confirmar la entrega actualiza Ventas e Inventario "
            "y conserva evidencia para consulta.",
        },
        {
            "title": "15. Finanzas",
            "purpose": "Controla cuentas por cobrar y pagar, cobros, pagos, presupuestos, centros de costo y conciliaciones.",
            "before": "Registre monedas, clientes, proveedores y centros de costo; facture ventas o compras para generar cuentas enlazadas.",
            "steps": [
                ("Revise las cuentas", "consulte vencimiento, saldo, moneda y documento de origen."),
                ("Registre un cobro", "seleccione la cuenta por cobrar, importe, fecha, método y referencia."),
                ("Registre un pago", "seleccione la cuenta por pagar y capture los datos del egreso."),
                ("Controle presupuesto", "cree el monto anual o mensual por centro de costo y categoría."),
                ("Concilie", "compare el saldo esperado con el saldo capturado y documente diferencias."),
                ("Verifique el resultado", "confirme saldos pendientes, vencidos y movimientos del periodo."),
            ],
            "controls": [
                "Las facturas de Ventas generan cuentas por cobrar.",
                "Las facturas de proveedor generan cuentas por pagar.",
                "Pagos y cobros reducen el saldo de la cuenta relacionada.",
                "Los presupuestos ayudan a contrastar autorización y compromiso.",
                "Las conciliaciones conservan diferencia y estado de revisión.",
            ],
            "connections": "Se alimenta de Compras y Ventas y puede relacionar gastos con centros de costo. Las decisiones y "
            "aprobaciones financieras pueden gestionarse en Tareas.",
        },
        {
            "title": "16. Tareas y aprobaciones",
            "purpose": "Organiza responsables, fechas límite, comentarios, decisiones, rechazos y reasignaciones.",
            "before": "Deben existir usuarios activos y, para flujos formales, permisos de validación o administración.",
            "steps": [
                ("Cree la tarea", "asigne título, responsable, prioridad, fecha límite y documento relacionado."),
                ("Defina el flujo", "establezca pasos y nivel requerido cuando la operación necesite aprobación."),
                ("Dé seguimiento", "use comentarios para conservar contexto y evidencia."),
                ("Decida", "apruebe o rechace cuando la tarea llegue al responsable autorizado."),
                ("Reasigne", "cambie al responsable sin perder el historial."),
                ("Cierre", "marque la tarea terminada y consulte el historial de decisiones."),
            ],
            "controls": [
                "Las fechas límite permiten identificar pendientes vencidos.",
                "Los comentarios quedan vinculados con la tarea.",
                "Los rechazos deben conservar el motivo.",
                "Las reasignaciones no eliminan las decisiones previas.",
                "El historial muestra responsable, acción y momento de cada cambio.",
            ],
            "connections": "Puede utilizarse para acompañar aprobaciones de compras, producción, calidad, mantenimiento, "
            "finanzas u otros documentos que requieran seguimiento humano.",
        },
    ]

    for module in modules:
        add_module(doc, module)

    # 17
    add_section_title(
        doc,
        "17. Flujo integral entre módulos",
        "Ejemplo de operación completa desde el pedido hasta la entrega.",
    )
    add_steps(
        doc,
        [
            ("Ventas confirma el pedido", "se establece producto, cantidad, fecha y cliente."),
            ("Producción revisa ingeniería", "consulta lista de materiales, ruta y capacidad."),
            ("Inventario informa disponibilidad", "reserva material disponible e identifica faltantes."),
            ("Compras atiende faltantes", "solicita, ordena y recibe material requerido."),
            ("Producción ejecuta", "consume materiales, registra tiempos y genera producto terminado."),
            ("Calidad inspecciona", "aprueba y libera, o rechaza para bloqueo o reproceso."),
            ("Inventario recibe", "registra el producto liberado como existencia disponible."),
            ("Logística prepara y entrega", "surte, empaca, documenta ruta y confirma al cliente."),
            ("Finanzas controla el cobro", "da seguimiento a la factura y registra el ingreso."),
        ],
    )
    add_note(
        doc,
        "Regla de trazabilidad",
        "No duplique documentos para acelerar el flujo. Siempre use la acción que parte del documento de origen; "
        "así se conservan cliente, artículos, cantidades, folios y estados relacionados.",
        "warning",
    )

    # 18
    add_section_title(
        doc,
        "18. Buenas prácticas operativas",
        "Hábitos recomendados para mantener datos confiables y reducir correcciones.",
    )
    add_bullets(
        doc,
        [
            ("Preparar catálogos", "cree maestros y catálogos antes de registrar documentos operativos."),
            ("Usar el documento origen", "genere pedidos desde cotizaciones, órdenes desde solicitudes y embarques desde pedidos."),
            ("No capturar folios", "permita que el sistema los asigne al guardar."),
            ("Revisar estados", "confirme que el documento esté aprobado o confirmado antes de la siguiente etapa."),
            ("Evitar duplicados", "busque por nombre, código o folio antes de crear un registro."),
            ("Registrar el momento real", "capture recepciones, consumos, paros, entregas, pagos y cobros cuando ocurren."),
            ("Documentar excepciones", "use notas, motivos, evidencias y comentarios para rechazos o diferencias."),
            ("Cerrar sesiones", "salga de la aplicación al terminar el turno."),
            ("Respaldar la base", "realice copias periódicas cuando el servidor esté detenido."),
        ],
    )
    doc.add_heading("Cierre diario sugerido", level=2)
    add_steps(
        doc,
        [
            ("Revise notificaciones", "atienda avisos pendientes y tareas próximas a vencer."),
            ("Revise documentos abiertos", "identifique solicitudes, órdenes o embarques sin avance."),
            ("Compare inventario crítico", "valide faltantes y reservas prioritarias."),
            ("Revise finanzas", "identifique cuentas vencidas, pagos y cobros no conciliados."),
            ("Consulte la bitácora", "verifique movimientos extraordinarios del día."),
        ],
    )

    # 19
    add_section_title(
        doc,
        "19. Solución de problemas",
        "Diagnóstico rápido de las incidencias más comunes.",
    )
    add_table(
        doc,
        ["Síntoma", "Causa probable", "Acción recomendada"],
        [
            ("Failed to fetch", "El servidor no está ejecutándose o se abrió solo Live Server.", "Ejecute npm start y confirme http://127.0.0.1:5050."),
            ("La página aparece sin estilos", "Ruta incorrecta o archivo CSS no disponible.", "Abra desde el servidor del ERP y actualice con Ctrl + F5."),
            ("No inicia sesión", "Credenciales incorrectas, cuenta bloqueada o servidor detenido.", "Revise el servidor y solicite al administrador validar la cuenta."),
            ("No aparece un artículo", "Está inactivo o no está habilitado para compra, venta, inventario o producción.", "Edite el artículo en Datos maestros y active la función necesaria."),
            ("No aparece un botón", "El rol no tiene permiso o el documento está en un estado incompatible.", "Revise rol, nivel, permiso y estado."),
            ("La pantalla regresa a otra opción", "Se utilizó una ruta anterior o la sesión perdió vigencia.", "Actualice la página y entre desde el menú lateral vigente."),
            ("No permite una salida", "La cantidad disponible es insuficiente o está reservada.", "Revise físico, reservado y disponible antes de retirar."),
            ("El folio muestra 000000", "Es una vista previa antes de guardar.", "Complete el formulario; el consecutivo real se asigna al guardar."),
            ("Fecha u hora incorrecta", "Zona horaria distinta a la operación.", "Cambie la lista de Zona horaria en Configuración."),
        ],
        [2200, 3200, 3960],
    )
    add_note(
        doc,
        "Antes de repetir una operación",
        "Si el navegador mostró un error después de guardar, busque primero el folio en la lista o la bitácora. "
        "Repetir la captura sin verificar puede generar un documento duplicado.",
        "warning",
    )

    # 20
    add_section_title(
        doc,
        "20. Respaldo y administración local",
        "La información se guarda en data/abicorp-erp.db dentro del proyecto, salvo que la instalación use otra ruta.",
    )
    doc.add_heading("Respaldo seguro", level=2)
    add_steps(
        doc,
        [
            ("Avise a los usuarios", "asegure que nadie esté capturando información."),
            ("Detenga el servidor", "cierre la ventana donde se ejecuta npm start."),
            ("Copie la base", "respalde data/abicorp-erp.db en una ubicación protegida y con fecha."),
            ("Respalde archivos", "si utiliza documentos adjuntos, copie también el directorio de datos asociado."),
            ("Reinicie y verifique", "ejecute npm start, inicie sesión y confirme el estado del sistema."),
        ],
    )
    add_note(
        doc,
        "Importante",
        "No sustituya ni copie la base mientras el servidor está escribiendo. Conserve varias generaciones de respaldo "
        "y proteja las copias porque contienen información operativa.",
        "warning",
    )
    doc.add_heading("Restablecer una cuenta administrativa", level=2)
    p = doc.add_paragraph()
    r = p.add_run("Desde la carpeta del proyecto, un administrador técnico puede ejecutar:\n")
    set_font(r)
    r = p.add_run("npm run reset-admin -- <usuario> <contraseña-segura>")
    set_font(r, name="Consolas", size=10, color=NAVY, bold=True)
    add_note(
        doc,
        "Seguridad",
        "No escriba contraseñas reales dentro de este manual ni las comparta por mensajes. Cambie cualquier "
        "contraseña temporal al primer acceso.",
    )

    # 21
    add_section_title(
        doc,
        "21. Referencia rápida",
        "Conceptos que aparecen de forma recurrente en la aplicación.",
    )
    add_table(
        doc,
        ["Concepto", "Significado"],
        [
            ("Folio", "Identificador consecutivo generado automáticamente para un documento."),
            ("Documento origen", "Registro previo desde el que se genera la siguiente etapa."),
            ("Borrador", "Registro creado que aún puede requerir revisión o aprobación."),
            ("Aprobado / confirmado", "Documento autorizado para avanzar al siguiente proceso."),
            ("En proceso", "Operación actualmente en ejecución."),
            ("Parcial", "Documento aplicado solo por una parte de la cantidad prevista."),
            ("Cerrado", "Proceso concluido; normalmente ya no admite nuevas acciones."),
            ("Cancelado", "Documento detenido sin continuar el flujo normal."),
            ("Existencia física", "Cantidad total registrada en un almacén."),
            ("Reservado", "Cantidad separada para una necesidad específica."),
            ("Disponible", "Existencia física menos cantidad reservada."),
            ("Bitácora", "Historial de eventos con usuario, fecha, módulo y descripción."),
        ],
        [2300, 7060],
    )
    doc.add_heading("Atajos de operación", level=2)
    add_bullets(
        doc,
        [
            "Ctrl + F5: recarga completa de la interfaz y sus estilos.",
            "Tab / Shift + Tab: avanza o retrocede entre controles del formulario.",
            "Enter: confirma botones y opciones enfocadas cuando el navegador lo permite.",
            "Escape: cierra algunas ventanas modales cuando no existe una restricción activa.",
            "Flechas izquierda y derecha: recorren el carrusel de módulos en la portada de acceso.",
        ],
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    r = p.add_run("FIN DEL MANUAL")
    set_font(r, size=10, color=TEAL, bold=True)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
